"""
extraer_matriz.py - Convierte la matriz de casos de prueba (Excel) de un
proyecto Web LDS a JSON, para que generate_scripts.py no tenga que leer el
.xlsx directamente (evita depender de openpyxl en la maquina de destino,
y evita re-transcribir a mano RF/CU/nombre/descripcion/etc. de cada caso).

Validado contra la matriz real de PR000335
("MCP-PR000335-Sprint1 (1).xlsx", hoja "Matriz de casos de prueba"):
encabezados en la fila 4, datos desde la fila 5. Columnas esperadas
(deteccion por NOMBRE de columna, no por letra, asi tolera que se agregue/
mueva alguna columna en otro proyecto):
    "Nro. RF asocidado"    -> rf
    "Nro. CU asociado"     -> cu
    "Nro. CP"               -> cp (clave del dict de salida, ej. "CP001")
    "Nombre CP"             -> name
    "Descripción del CP"    -> desc
    "Precondición"          -> precond
    "Datos de pruebas"      -> datos
    "Resultado esperado"    -> esperado
    "Pasos a seguir"        -> pasos

Requiere openpyxl (`pip install openpyxl`) SOLO para correr este script -
no es una dependencia de los CPxxx.py generados ni de run_all.py.

Uso:
    .venv\\Scripts\\python.exe extraer_matriz.py "ruta\\a\\Matriz.xlsx" salida.json
    .venv\\Scripts\\python.exe extraer_matriz.py "ruta\\a\\Matriz.xlsx" salida.json --hoja "Matriz de casos de prueba"

Con la salida (salida.json) el nuevo generate_scripts.py puede:
    - Tomar directamente "pasos" para el docstring de cada CP (como ya hace
      pasos_matriz.json en PR000335).
    - Usar rf/cu/name/desc/precond/datos/esperado como fuente de verdad
      para llamar a add(...) sin re-tipear cada campo a mano (menos
      transcripcion manual = menos errores respecto al Excel real).
"""
import argparse
import json
import sys
from pathlib import Path

# Nombres de columna esperados (encabezado exacto) -> nombre de campo en
# el JSON de salida. Ajustar aca si otro proyecto usa encabezados distintos.
COLUMNAS = {
    "Nro. RF asocidado": "rf",
    "Nro. RF asociado": "rf",  # variante sin el typo, por si se corrige en otro proyecto
    "Nro. CU asociado": "cu",
    "Nro. CP": "cp",
    "Nombre CP": "name",
    "Descripción del CP": "desc",
    "Precondición": "precond",
    "Datos de pruebas": "datos",
    "Resultado esperado": "esperado",
    "Pasos a seguir": "pasos",
}


def encontrar_fila_encabezado(ws, max_filas_buscar=15):
    """Busca la primera fila que contenga "Nro. CP" (o similar) en alguna
    celda: la matriz real trae 2-3 filas de titulo/nota arriba del
    encabezado real (confirmado en PR000335: encabezado en la fila 4, no
    la 1)."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_filas_buscar, values_only=True), start=1):
        for celda in row:
            if isinstance(celda, str) and celda.strip().startswith("Nro. CP"):
                return i
    raise SystemExit(
        f"No se encontro una fila de encabezado con 'Nro. CP' en las primeras {max_filas_buscar} filas. "
        "Revisar el Excel a mano o ajustar max_filas_buscar."
    )


def extraer(ruta_excel: Path, hoja: str | None = None) -> dict:
    import openpyxl  # import diferido: solo hace falta para correr este script

    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    nombre_hoja = hoja or next(
        (n for n in wb.sheetnames if "matriz" in n.lower() and "caso" in n.lower()),
        wb.sheetnames[0],
    )
    if nombre_hoja not in wb.sheetnames:
        raise SystemExit(f"Hoja '{nombre_hoja}' no existe. Hojas disponibles: {wb.sheetnames}")
    ws = wb[nombre_hoja]

    fila_encabezado = encontrar_fila_encabezado(ws)
    encabezados = [c.value for c in ws[fila_encabezado]]
    columna_por_campo = {}
    for idx, texto in enumerate(encabezados):
        if isinstance(texto, str):
            campo = COLUMNAS.get(texto.strip())
            if campo:
                columna_por_campo[campo] = idx

    faltantes = [c for c in ("cp", "pasos") if c not in columna_por_campo]
    if faltantes:
        raise SystemExit(
            f"No se encontraron las columnas obligatorias {faltantes} en la fila de encabezado "
            f"({fila_encabezado}). Encabezados leidos: {encabezados}. "
            "Ajustar el diccionario COLUMNAS de este script si el nombre exacto difiere."
        )

    resultado = {}
    for row in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        cp = row[columna_por_campo["cp"]] if columna_por_campo.get("cp") is not None else None
        if not cp or not isinstance(cp, str) or not cp.strip().upper().startswith("CP"):
            continue  # fila vacia o de cierre de tabla
        cp = cp.strip()
        caso = {}
        for campo, idx in columna_por_campo.items():
            valor = row[idx] if idx < len(row) else None
            caso[campo] = (valor or "").strip() if isinstance(valor, str) else (valor or "")
        resultado[cp] = caso

    return resultado


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("excel", help="Ruta al .xlsx de la matriz de casos de prueba")
    p.add_argument("salida", help="Ruta del .json a generar")
    p.add_argument("--hoja", default=None, help="Nombre exacto de la hoja (default: autodetecta la que tenga 'matriz' y 'caso' en el nombre)")
    p.add_argument(
        "--solo-pasos",
        action="store_true",
        help="Generar SOLO {cp: pasos} en vez de todos los campos (compatible con pasos_matriz.json de PR000335)",
    )
    args = p.parse_args()

    datos = extraer(Path(args.excel), args.hoja)
    if args.solo_pasos:
        datos = {cp: caso.get("pasos", "") for cp, caso in datos.items()}

    Path(args.salida).write_text(
        json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Extraidos {len(datos)} casos -> {args.salida}")


if __name__ == "__main__":
    main()
