"""
agregar_hallazgos_reporte.py (especifico de PR000349-SP2, NO generico).

Agrega una hoja "Hallazgos para Desarrollo" al reporte consolidado mas
reciente generado por generar_reporte_excel.py - narrativa estructurada
de los defectos/bloqueos reales confirmados en vivo esta sesion, para
que Desarrollo pueda actuar sobre ellos (no solo ver el semaforo OK/
FALLIDO/SIN CORRER de la hoja "Detalle por caso").

Uso: correr DESPUES de generar_reporte_excel.py.
    .venv\\Scripts\\python.exe agregar_hallazgos_reporte.py
"""
import glob
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
REPORTES_DIR = BASE / "reportes"

AZUL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FUENTE_HEADER = Font(color="FFFFFF", bold=True)
ROJO = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
NARANJA = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
AMARILLO = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
GRIS_CLARO = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
FUENTE_SEV = Font(bold=True)

FILL_POR_SEVERIDAD = {
    "CRÍTICO": ROJO,
    "ALTO": NARANJA,
    "MEDIO": AMARILLO,
    "OBSERVACIÓN": GRIS_CLARO,
}

HALLAZGOS = [
    {
        "n": 1,
        "severidad": "CRÍTICO",
        "rf": "RF-07, RF-09 (posiblemente RF-03/05/08 también, sin confirmar aún)",
        "titulo": "El componente de fecha (nz-date-picker) congela la pantalla por completo",
        "descripcion": (
            "Al ingresar una fecha calendario VÁLIDA en un filtro que usa el "
            "componente nz-date-picker (ng-zorro, nzformat=\"dd/MM/yyyy\"), la "
            "pantalla completa deja de responder a cualquier interacción "
            "(clicks, lectura de campos, capturas). Reproducido de forma "
            "independiente en 2 pantallas Angular distintas (Consulta de "
            "Cortes Pendientes y Consulta Genérica de Solicitudes de "
            "Reposición) con 5+ métodos de interacción distintos (tipeo "
            "normal, tipeo lento 300ms/caracter, .fill(), click en el "
            "calendario visual, click en botón \"Hoy\"). NO se recupera ni "
            "esperando 100+ segundos. Se descartó: velocidad de "
            "automatización (probado con delay), diálogo nativo bloqueante "
            "(page.on(\"dialog\") no detectó ninguno), excepción de "
            "JavaScript (0 errores de consola). Una fecha INVÁLIDA o el "
            "campo vacío NO producen el freeze - es específico de un valor "
            "de fecha calendario real aceptado por el componente."
        ),
        "mensaje_real": "Ninguno - la pantalla simplemente deja de responder (sin diálogo, sin error de consola)",
        "endpoint_o_componente": "nz-date-picker (ng-zorro/Ant Design)",
        "casos_bloqueados": "CP052, CP053 (RF-07); CP064, CP068, CP069, CP070, CP071 (RF-09) — 7 casos",
        "evidencia": "_exploracion/explore_21..27_rf07_*.py, explore_28..31_rf09_*.py, explore_38_rf07_espera_larga.py",
    },
    {
        "n": 2,
        "severidad": "ALTO",
        "rf": "RF-02 Mantenimiento de Órdenes de Conexión",
        "titulo": "Backend caído: endpoint de búsqueda responde 503",
        "descripcion": (
            "El botón \"Leer\" (búsqueda individual de cliente) llama al "
            "endpoint mantord/Leer, que responde HTTP 503 para CUALQUIER "
            "número de cliente, no solo para casos puntuales. Confirmado "
            "con 3 clientes reales distintos (264252, 781409, 1201729), "
            "3/3 con el mismo resultado. La pantalla muestra "
            "\"Error de Servicio: Ocurrió un error inesperado. Intente "
            "nuevamente.\" Bloquea la pantalla completa, no es un problema "
            "de datos de prueba."
        ),
        "mensaje_real": "Error de Servicio: Ocurrió un error inesperado. Intente nuevamente. (HTTP 503)",
        "endpoint_o_componente": "https://tk0s-web.luzdelsur.com.pe/comercial/gesclie/acticome/nucli/web/api/mantord/Leer",
        "casos_bloqueados": "CP009-CP016, CP072 — 9 casos",
        "evidencia": "_exploracion/explore_62_rf02_verificar_503.py",
    },
    {
        "n": 3,
        "severidad": "ALTO",
        "rf": "RF-03 Consulta de Cortes y Reconexiones",
        "titulo": "Backend caído: endpoint de búsqueda responde 504 (Gateway Timeout)",
        "descripcion": (
            "El botón \"Leer\" con un filtro válido (Sucursal Chacarilla, "
            "sin fechas) llama al endpoint BuscarCortesReconexiones, que "
            "responde HTTP 504. La pantalla muestra "
            "\"Error de Servicio: Ocurrió un error inesperado. Intente "
            "nuevamente.\" Reportado primero por el usuario en su propia "
            "sesión manual, reproducido de forma independiente en "
            "automatización."
        ),
        "mensaje_real": "Error de Servicio: Ocurrió un error inesperado. Intente nuevamente. (HTTP 504)",
        "endpoint_o_componente": ".../corterepos/corterepos/corterepos/web/api/concrcli/BuscarCortesReconexiones",
        "casos_bloqueados": "CP017-CP024 — 8 casos",
        "evidencia": "_exploracion/explore_42_rf03_leer_504.py",
    },
    {
        "n": 4,
        "severidad": "ALTO",
        "rf": "RF-05 Consulta de Pagos",
        "titulo": "Backend caído: la pantalla no llega a cargar (503)",
        "descripcion": (
            "El endpoint de carga inicial CargarPantalla "
            "(recauda/pagoscent/consultapago) responde HTTP 503 de forma "
            "consistente (4/4 intentos) - la pantalla ni siquiera llega a "
            "mostrarse. Control cruzado confirmado: en el mismo momento, "
            "otras pantallas del mismo sprint (RF-07, RF-09) sí cargaban "
            "bien - no es un problema de ambiente general."
        ),
        "mensaje_real": "HTTP 503 en CargarPantalla (sin llegar a renderizar la pantalla)",
        "endpoint_o_componente": ".../recauda/pagoscent/consultapago (CargarPantalla)",
        "casos_bloqueados": "CP031-CP037 — 7 casos",
        "evidencia": "Confirmado 4/4 intentos, sesión 27-28/09/2026",
    },
    {
        "n": 5,
        "severidad": "ALTO",
        "rf": "RF-08 Consulta de Solicitudes de Reposición Pendientes",
        "titulo": "Falla parcial: el combo de Sucursal Cliente no carga",
        "descripcion": (
            "El resto de la pantalla carga correctamente, pero el combo "
            "\"Sucursal Cliente\" falla con \"No se pudo cargar el combo de "
            "Sucursal Cliente\" (3/3 intentos) - un filtro clave queda "
            "inutilizable."
        ),
        "mensaje_real": "No se pudo cargar el combo de Sucursal Cliente",
        "endpoint_o_componente": "Combo Sucursal Cliente (RF-08)",
        "casos_bloqueados": "CP057-CP063 — 7 casos",
        "evidencia": "Confirmado 3/3 intentos, sesión 27-28/09/2026",
    },
    {
        "n": 6,
        "severidad": "MEDIO",
        "rf": "RF-06 Mantenedor de Inspectores de Corte",
        "titulo": "Alta de inspector bloqueada: servicio de validación de ficha caído",
        "descripcion": (
            "CUALQUIER intento de dar de Alta un inspector falla con "
            "\"Error al validar la ficha.\" ANTES de evaluar la regla de "
            "negocio del caso - confirmado 2/2 veces, tanto con una ficha "
            "inventada (999001) como con una ficha REAL ya asignada (4504, "
            "la del inspector 1381). Coincide exactamente con un riesgo YA "
            "documentado en el EF (Sección 11): el Alta depende de un "
            "servicio compartido de generación de correlativo y resolución "
            "de área por sucursal, cuya disponibilidad es prerrequisito de "
            "cierre. Efecto colateral: al no poder crear inspectores de "
            "prueba descartables, CP044/045/048 no quedaron como casos "
            "automatizados repetibles (se verificaron manualmente una vez, "
            "con evidencia real, para no arriesgar escrituras repetidas "
            "sobre datos reales)."
        ),
        "mensaje_real": "Error de Servicio: Error al validar la ficha.",
        "endpoint_o_componente": "Servicio de catálogos compartido (correlativo + área por sucursal) — RF-06 Alta",
        "casos_bloqueados": "CP046, CP047 directos; CP044, CP045, CP048 indirectos (sin datos descartables renovables) — 5 casos",
        "evidencia": "_exploracion/explore_51_rf06_alta_todas_sucursales.py, explore_58_rf06_cp046_solo.py",
    },
    {
        "n": 7,
        "severidad": "MEDIO",
        "rf": "RF-01 Consulta de Clientes Asociados a una S.S.",
        "titulo": "Generación de informe (Imprimir) no disponible",
        "descripcion": (
            "La generación del informe falla consistentemente con "
            "\"No se generó informe.\" (confirmado 3 veces, no transitorio) "
            "- coincide con un riesgo YA documentado en el EF (Sección 11): "
            "el servicio de generación de reportes es externo al alcance "
            "de este sprint, su disponibilidad es un prerrequisito de "
            "cierre. El diálogo de confirmación previo funciona bien."
        ),
        "mensaje_real": "No se generó informe.",
        "endpoint_o_componente": "Servicio externo de generación de reportes — RF-01",
        "casos_bloqueados": "CP008 — 1 caso",
        "evidencia": "run_all.py --solo CP008 (3 corridas)",
    },
    {
        "n": 8,
        "severidad": "OBSERVACIÓN",
        "rf": "RF-06 Mantenedor de Inspectores de Corte",
        "titulo": "El universo de inspectores activos en Pruebas supera el límite de 2000 registros",
        "descripcion": (
            "La matriz de casos asume que ejecutar \"Leer\" con el filtro "
            "por defecto (Estado=Activos, sin ningún otro filtro) devuelve "
            "un resultado dentro del límite soportado por la grilla (2000 "
            "registros). En la práctica, el ambiente de Pruebas tiene MÁS "
            "de 2000 inspectores activos (y también más de 2000 inactivos), "
            "por lo que esa combinación siempre da \"Demasiados registros "
            "para ser mostrados en la grilla.\" (comportamiento correcto de "
            "RN-03, no un defecto) - la matriz necesitaría un filtro "
            "adicional para probar el caso de \"resultado dentro del "
            "límite\" con datos reales de Pruebas."
        ),
        "mensaje_real": "Demasiados registros para ser mostrados en la grilla.",
        "endpoint_o_componente": "RF-06 — filtro por defecto",
        "casos_bloqueados": "CP039 — 1 caso",
        "evidencia": "_exploracion/explore_45_rf06_completo.py",
    },
    {
        "n": 9,
        "severidad": "OBSERVACIÓN",
        "rf": "RF-06 Mantenedor de Inspectores de Corte",
        "titulo": "Filas duplicadas en la grilla al filtrar por Código (coincidencia parcial)",
        "descripcion": (
            "Al filtrar por Código=\"1\" (coincidencia parcial), varios "
            "códigos reales (1378/1379/1380/1381) devuelven CADA UNO "
            "múltiples filas idénticas en la grilla (mismos valores en "
            "todas las columnas) - no son inspectores distintos. No se "
            "concluye como defecto (podría ser data real así en Pruebas); "
            "queda anotado para que Desarrollo lo revise si le resulta "
            "extraño."
        ),
        "mensaje_real": "(sin mensaje - duplicación silenciosa de filas)",
        "endpoint_o_componente": "RF-06 — grilla de resultados",
        "casos_bloqueados": "Ninguno directamente (observación general)",
        "evidencia": "_exploracion/explore_47_rf06_todos_codigo1.py",
    },
]


def agregar_hoja_hallazgos(ruta_xlsx: Path):
    wb = load_workbook(ruta_xlsx)
    if "Hallazgos para Desarrollo" in wb.sheetnames:
        del wb["Hallazgos para Desarrollo"]
    ws = wb.create_sheet("Hallazgos para Desarrollo", 0)

    ws.append(["Hallazgos confirmados en vivo — PR000349-SP2 — para reportar a Desarrollo"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{len(HALLAZGOS)} hallazgos: 1 crítico, 4 altos (backends caídos), 2 medios, 2 observaciones"])
    ws.append([])

    encabezados = ["#", "Severidad", "RF afectado(s)", "Título", "Descripción",
                   "Mensaje real observado", "Endpoint / Componente",
                   "Casos de prueba bloqueados", "Evidencia (script de exploración)"]
    ws.append(encabezados)
    fila_header = ws.max_row
    for c in ws[fila_header]:
        c.font = FUENTE_HEADER
        c.fill = AZUL_HEADER

    for h in HALLAZGOS:
        ws.append([
            h["n"], h["severidad"], h["rf"], h["titulo"], h["descripcion"],
            h["mensaje_real"], h["endpoint_o_componente"], h["casos_bloqueados"], h["evidencia"],
        ])
        fila = ws.max_row
        celda_sev = ws.cell(row=fila, column=2)
        celda_sev.fill = FILL_POR_SEVERIDAD.get(h["severidad"], GRIS_CLARO)
        celda_sev.font = FUENTE_SEV
        celda_sev.alignment = Alignment(horizontal="center", vertical="top")
        for col in range(1, len(encabezados) + 1):
            celda = ws.cell(row=fila, column=col)
            celda.border = BORDE
            celda.alignment = Alignment(vertical="top", wrap_text=(col in (4, 5, 6, 7, 8, 9)))

    anchos = [4, 12, 22, 30, 60, 35, 35, 35, 45]
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = "A5"

    wb.save(ruta_xlsx)
    print(f"Hoja 'Hallazgos para Desarrollo' agregada a: {ruta_xlsx}")


def main():
    candidatos = sorted(REPORTES_DIR.glob("reporte_consolidado_*.xlsx"))
    if not candidatos:
        raise SystemExit("No se encontro ningun reporte_consolidado_*.xlsx - correr generar_reporte_excel.py primero.")
    mas_reciente = candidatos[-1]
    agregar_hoja_hallazgos(mas_reciente)


if __name__ == "__main__":
    main()
